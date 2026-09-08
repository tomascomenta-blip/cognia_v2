# -*- coding: utf-8 -*-
"""
tests/test_cotidiano_tools.py
=============================
Tareas cotidianas en el escritorio propio (cognia/agent/cotidiano_tools.py,
2026-09-08). Sin ventanas ni cuentas reales: la config de correo en config.env,
el envio por SMTP contra un servidor local de verdad (con adjunto), el fallo
HONESTO sin cuenta (dice como configurarla, no finge), los documentos (.txt,
.md y .docx leidos de vuelta), las fechas y la cita .ics. Con
COGNIA_E2E_APP=1 ademas abre el Bloc de notas en el escritorio de Cognia.
"""
from __future__ import annotations

import asyncio
import os
import socket
import threading
from pathlib import Path

import pytest

from cognia.agent import cotidiano_tools as CT


@pytest.fixture(autouse=True)
def _aislar(tmp_path, monkeypatch):
    monkeypatch.setenv("COGNIA_HOME", str(tmp_path / "home"))
    for k in ("CORREO_SMTP_HOST", "CORREO_SMTP_PORT", "CORREO_USUARIO", "CORREO_CLAVE", "CORREO_DE",
              "CORREO_IMAP_HOST", "CORREO_TLS"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setattr(CT, "outlook_tiene_cuenta", lambda: False)


# ── config ──────────────────────────────────────────────────────────────────

def test_config_correo_vacia_y_guardada(tmp_path):
    assert CT.config_correo()["host"] == ""
    r = CT.guardar_config_correo("yo@gmail.com", "clave secreta")
    assert r == tmp_path / "home" / "config.env"
    c = CT.config_correo()
    assert c["host"] == "smtp.gmail.com" and c["imap"] == "imap.gmail.com" and c["clave"] == "clave secreta"
    # regrabar no duplica claves
    CT.guardar_config_correo("yo@gmail.com", "otra", host="smtp.x.com", puerto="465")
    texto = r.read_text(encoding="utf-8")
    assert texto.count("CORREO_CLAVE=") == 1 and "CORREO_SMTP_HOST=smtp.x.com" in texto


def test_enviar_sin_cuenta_dice_que_falta():
    with pytest.raises(ValueError) as e:
        CT.enviar_correo("a@b.com", "hola", "cuerpo")
    assert "correo_configurar" in str(e.value) and "Outlook: sin cuenta" in str(e.value)
    with pytest.raises(ValueError, match="destinatario"):
        CT.enviar_correo("sin-arroba", "hola", "cuerpo")


# ── SMTP real contra un servidor local ─────────────────────────────────────

class _SMTPLocal:
    """Servidor SMTP minimo en un hilo: guarda el ultimo mensaje."""

    def __init__(self):
        s = socket.socket()
        s.bind(("127.0.0.1", 0))
        self.puerto = s.getsockname()[1]
        s.close()
        self.mensajes = []
        self._loop = None
        self._hilo = threading.Thread(target=self._correr, daemon=True)
        self._listo = threading.Event()

    async def _atender(self, r, w):
        w.write(b"220 local ESMTP\r\n")
        await w.drain()
        datos, en_data = b"", False
        while True:
            linea = await r.readline()
            if not linea:
                break
            if en_data:
                if linea.strip() == b".":
                    en_data = False
                    self.mensajes.append(datos)
                    w.write(b"250 OK\r\n")
                else:
                    datos += linea
                await w.drain()
                continue
            cmd = linea.strip().upper()
            if cmd.startswith((b"EHLO", b"HELO")):
                w.write(b"250-local\r\n250 8BITMIME\r\n")
            elif cmd.startswith(b"DATA"):
                en_data = True
                w.write(b"354 go\r\n")
            elif cmd.startswith(b"QUIT"):
                w.write(b"221 bye\r\n")
                await w.drain()
                w.close()
                return
            else:
                w.write(b"250 OK\r\n")
            await w.drain()

    def _correr(self):
        self._loop = asyncio.new_event_loop()
        asyncio.set_event_loop(self._loop)

        async def main():
            srv = await asyncio.start_server(self._atender, "127.0.0.1", self.puerto)
            self._listo.set()
            async with srv:
                await srv.serve_forever()
        try:
            self._loop.run_until_complete(main())
        except Exception:
            pass

    def __enter__(self):
        self._hilo.start()
        self._listo.wait(5)
        return self

    def __exit__(self, *a):
        # se deja morir con el hilo daemon: parar el loop desde fuera dispara
        # el ResourceWarning del proactor de Windows sin aportar nada
        return False


def test_enviar_por_smtp_local_con_adjunto(tmp_path, monkeypatch):
    adj = tmp_path / "informe.txt"
    adj.write_text("informe de prueba", encoding="utf-8")
    with _SMTPLocal() as srv:
        monkeypatch.setenv("CORREO_SMTP_HOST", "127.0.0.1")
        monkeypatch.setenv("CORREO_SMTP_PORT", str(srv.puerto))
        monkeypatch.setenv("CORREO_TLS", "0")
        monkeypatch.setenv("CORREO_DE", "cognia@local")
        r = CT.enviar_correo("dueno@example.com, otro@example.com", "Prueba", "Hola\n\nCognia",
                             cc="copia@example.com", adjunto=str(adj))
        assert "enviado por SMTP" in r and "dueno@example.com" in r
        import time
        for _ in range(20):
            if srv.mensajes:
                break
            time.sleep(0.1)
        assert srv.mensajes, "el servidor local no recibio el mensaje"
        crudo = srv.mensajes[0].decode("utf-8", "replace")
        assert "Subject: Prueba" in crudo and "To: dueno@example.com, otro@example.com" in crudo
        assert "Cc: copia@example.com" in crudo and 'filename="informe.txt"' in crudo
        # la tool por run_tool
        from cognia.agent.tools import run_tool
        out = run_tool("correo_enviar", "dueno@example.com | asunto=Dos | cuerpo=linea1\\nlinea2", {})
        assert out.startswith("RESULTADO correo_enviar: enviado por SMTP")


# ── documentos ──────────────────────────────────────────────────────────────

def test_escribir_documento_txt_md_y_docx(tmp_path):
    r = CT.escribir_documento(tmp_path / "nota.txt", "hola\n\nadios", titulo="Nota")
    assert r["formato"] == "txt" and (tmp_path / "nota.txt").read_text(encoding="utf-8").startswith("Nota\n====")
    r = CT.escribir_documento(tmp_path / "nota.md", "hola", titulo="Nota")
    assert (tmp_path / "nota.md").read_text(encoding="utf-8").startswith("# Nota")
    pytest.importorskip("docx")
    r = CT.escribir_documento(tmp_path / "carta.docx", "Estimado dueno,\n\n- uno\n- dos\n\n# Cierre\n\nSaludos.", titulo="Carta")
    assert r["formato"] == "docx" and r["parrafos"] == 4
    import docx
    d = docx.Document(str(tmp_path / "carta.docx"))
    textos = [p.text for p in d.paragraphs if p.text.strip()]
    assert textos[0] == "Carta" and "Estimado dueno," in textos and "uno" in textos and "Cierre" in textos
    estilos = {p.style.name for p in d.paragraphs}
    assert "List Bullet" in estilos and "Heading 1" in estilos


def test_documento_escribir_por_run_tool_sin_abrir(tmp_path):
    from cognia.agent.tools import run_tool
    out = run_tool("documento_escribir", "carta.txt | texto=Hola\\nmundo | abrir=0", {"workspace": str(tmp_path)})
    assert out.startswith("RESULTADO documento_escribir") and (tmp_path / "carta.txt").read_text(encoding="utf-8") == "Hola\nmundo\n"
    out = run_tool("documento_escribir", "carta2.txt | abrir=0", {"workspace": str(tmp_path)})
    assert "ERROR" in out and "texto=" in out


# ── fechas, citas y recordatorios ───────────────────────────────────────────

def test_fechas_y_cuando():
    import datetime as dt
    assert CT._parsear_fecha("2026-09-08 10:30") == dt.datetime(2026, 9, 8, 10, 30)
    assert CT._parsear_fecha("manana 09:15").hour == 9
    with pytest.raises(ValueError, match="fecha"):
        CT._parsear_fecha("el martes")
    ahora = dt.datetime.now()
    assert (CT._cuando(en="20m") - ahora).total_seconds() == pytest.approx(1200, abs=5)
    assert (CT._cuando(en="2h") - ahora).total_seconds() == pytest.approx(7200, abs=5)
    assert CT._cuando(a="23:59") > ahora
    with pytest.raises(ValueError):
        CT._cuando()


def test_cita_sin_outlook_deja_ics(tmp_path, monkeypatch):
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    monkeypatch.setenv("HOME", str(tmp_path))
    r = CT.agregar_cita("Dentista", "2026-10-01 16:00", duracion=30, lugar="Clinica")
    assert r["via"] == "ics"
    ics = Path(r["detalle"])
    assert ics.exists() and ics.parent == tmp_path / ".cognia" / "cotidiano" / "citas"
    t = ics.read_text(encoding="utf-8")
    assert "SUMMARY:Dentista" in t and "DTSTART:20261001T160000" in t and "DTEND:20261001T163000" in t and "LOCATION:Clinica" in t


def test_recordatorio_rechaza_lo_imposible():
    import datetime as dt
    with pytest.raises(ValueError, match="paso"):
        CT.programar_recordatorio("x", dt.datetime.now() - dt.timedelta(minutes=1))
    with pytest.raises(ValueError, match="texto"):
        CT.programar_recordatorio("   ", dt.datetime.now() + dt.timedelta(minutes=5))


# ── registro y estado ───────────────────────────────────────────────────────

def test_registro_y_estado():
    from cognia.agent.tools import TOOLS, run_tool
    for n in ("documento_escribir", "correo_enviar", "correo_leer", "correo_configurar", "calendario_agregar",
              "recordatorio", "abrir_en_escritorio", "cotidiano_estado"):
        assert n in TOOLS and CT.es_de_la_familia(n)
    assert TOOLS["correo_enviar"]["danger"] and not TOOLS["documento_escribir"]["danger"]
    out = run_tool("cotidiano_estado", "", {})
    assert "SMTP sin configurar" in out and "Outlook sin cuenta" in out


def test_familia_visible_por_flag(monkeypatch):
    from cognia.agent.tools import flag_de_optin
    assert flag_de_optin("correo_enviar") == "COGNIA_COTIDIANO"
    assert flag_de_optin("abrir_en_escritorio") == "COGNIA_COTIDIANO"
    from cognia.harness import familias
    fila = next(f for f in familias.estado() if f["familia"] == "cotidiano")
    assert fila["n_tools"] >= 8


@pytest.mark.skipif(not os.environ.get("COGNIA_E2E_APP"), reason="e2e real: COGNIA_E2E_APP=1 (Windows, pyvda)")
def test_e2e_bloc_de_notas_en_el_escritorio_propio(tmp_path):
    from cognia.agent.tools import run_tool
    from cognia.agent import app_tools as AT
    out = run_tool("documento_escribir", "carta.txt | titulo=Hola | texto=escrito por Cognia", {"workspace": str(tmp_path)})
    try:
        assert "abierto en el escritorio de Cognia" in out and "Bloc de notas" in out
    finally:
        AT.cerrar_todas()
