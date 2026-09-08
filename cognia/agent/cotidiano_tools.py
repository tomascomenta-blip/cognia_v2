# -*- coding: utf-8 -*-
"""
cognia/agent/cotidiano_tools.py
===============================
TAREAS COTIDIANAS en el escritorio propio de Cognia (2026-09-08). Pedido del
dueno: "que el segundo monitor sea mas funcional: que pueda hacer tareas
cotidianas que yo le pida, como escribir algo, mandar un correo, cosas asi".

El "segundo monitor" es el escritorio virtual 'Cognia' (escritorio_propio.py):
todo lo que se abre aqui se muda alli nada mas aparecer su ventana, asi que el
dueno lo ve cuando quiere (Win+Ctrl+flecha) y no le estorba mientras tanto.

Tools (prefijos correo_ / documento_ / calendario_ + abrir_en_escritorio,
recordatorio, cotidiano_estado; flag COGNIA_COTIDIANO, config `cotidiano_tools`,
ENCENDIDA por defecto):

  documento_escribir <ruta.docx|.txt|.md> | texto=... [| titulo=...] [| abrir=0]
      Escribe el documento (Word real por python-docx; texto plano/markdown tal
      cual) y lo ABRE en el escritorio de Cognia (Word para .docx, Bloc de notas
      para el resto). Devuelve la ruta, la ventana y una captura.
  correo_enviar <para> | asunto=... | cuerpo=... [| cc=...] [| adjunto=ruta] [| html=1]
      Manda el correo por el backend disponible, en este orden: SMTP configurado
      (Gmail con clave de aplicacion: `correo_configurar`), Outlook por COM si
      tiene una cuenta (en un subproceso con timeout: Outlook sin cuenta abre su
      asistente y bloquea COM para siempre, MEDIDO 2026-09-08). Sin ninguno,
      dice EXACTAMENTE que falta y como configurarlo.
  correo_leer [| n=10] [| buscar=texto] [| no_leidos=1]
      Lee la bandeja por IMAP (misma configuracion) u Outlook si hay cuenta.
  correo_configurar usuario=... clave=... [| host=smtp.gmail.com] [| puerto=587] [| imap=imap.gmail.com]
      Guarda la configuracion en ~/.cognia/config.env (nunca en el JSON de config).
  calendario_agregar <asunto> | inicio=YYYY-MM-DD HH:MM [| duracion=60] [| lugar=...] [| cuerpo=...]
      Cita en Outlook si hay cuenta; si no, un .ics en ~/.cognia/cotidiano/citas/ (importable).
  recordatorio <texto> | en=20m | a=HH:MM
      Aviso en pantalla a esa hora (tarea programada de Windows con msg.exe).
  abrir_en_escritorio <ruta|URL|app> [| titulo=...] [| espera=MS]
      Abre lo que sea y muda su ventana al escritorio de Cognia (URL: ventana
      NUEVA de Chrome/Edge con perfil propio, para no tocar el navegador del dueno).
  cotidiano_estado
      Que backend hay para cada cosa y el ultimo error.

Convenciones: las de la familia de pruebas (RESULTADO <tool> ...: / ERROR:;
args '<objetivo> | clave=valor'; rutas relativas contra scratchpad/workspace/cwd).
Todo fallo de backend se devuelve con causa; nada se calla.
"""
from __future__ import annotations

import datetime as _dt
import json
import os
import re
import subprocess
import sys
import time
from pathlib import Path

from cognia.agent import pruebas_comun as PC

_PREFIJOS = ("correo_", "documento_", "calendario_")
_NOMBRES = ("abrir_en_escritorio", "recordatorio", "cotidiano_estado")
_CLAVES = ("texto", "titulo", "abrir", "asunto", "cuerpo", "cc", "adjunto", "html", "n", "buscar",
           "no_leidos", "usuario", "clave", "host", "puerto", "imap", "de", "inicio", "duracion",
           "lugar", "en", "a", "espera", "tls", "borrar")
TIMEOUT_OUTLOOK_S = 45
ESPERA_VENTANA_MS = 8000

_ULTIMO: dict = {"tool": "", "detalle": "", "error": "", "ts": 0.0}


def es_de_la_familia(nombre: str) -> bool:
    return nombre.startswith(_PREFIJOS) or nombre in _NOMBRES


def _anotar(tool: str, detalle: str = "", error: str = "") -> None:
    _ULTIMO.update({"tool": tool, "detalle": str(detalle)[:300], "error": str(error)[:300], "ts": time.time()})


def _avisar(motivo: str) -> None:
    try:
        from cognia.cli import _aviso_degradado
        _aviso_degradado("cotidiano", motivo)
    except Exception:
        print("[degradado] cotidiano: %s" % motivo, file=sys.stderr)


def _err(tool: str, exc) -> str:
    _anotar(tool, "", str(exc))
    return "RESULTADO %s ERROR: %s" % (tool, exc)


def ultimo() -> dict:
    return dict(_ULTIMO)


# ---------------------------------------------------------------------------
# Config de correo: ~/.cognia/config.env (secretos fuera del JSON)
# ---------------------------------------------------------------------------

def _config_env_ruta() -> Path:
    crudo = os.environ.get("COGNIA_HOME", "").strip()
    base = Path(crudo) if crudo else Path.home() / ".cognia"
    return base / "config.env"


def _leer_config_env() -> dict:
    out = {}
    try:
        r = _config_env_ruta()
        if r.exists():
            for linea in r.read_text(encoding="utf-8", errors="replace").splitlines():
                linea = linea.strip()
                if not linea or linea.startswith("#") or "=" not in linea:
                    continue
                k, v = linea.split("=", 1)
                out[k.strip()] = v.strip().strip('"')
    except Exception as exc:
        _avisar("config.env ilegible: %s" % exc)
    return out


def config_correo() -> dict:
    """{host, puerto, usuario, clave, de, imap, tls} de env vars o config.env. Vacio = sin configurar."""
    env = _leer_config_env()

    def g(k, d=""):
        return os.environ.get(k, "").strip() or env.get(k, d)
    c = {"host": g("CORREO_SMTP_HOST", ""), "puerto": g("CORREO_SMTP_PORT", "587"),
         "usuario": g("CORREO_USUARIO", ""), "clave": g("CORREO_CLAVE", ""),
         "de": g("CORREO_DE", "") or g("CORREO_USUARIO", ""), "imap": g("CORREO_IMAP_HOST", ""),
         "tls": g("CORREO_TLS", "1")}
    if c["usuario"] and not c["host"] and c["usuario"].lower().endswith("@gmail.com"):
        c["host"] = "smtp.gmail.com"
    if c["usuario"] and not c["imap"] and c["usuario"].lower().endswith("@gmail.com"):
        c["imap"] = "imap.gmail.com"
    return c


def guardar_config_correo(usuario: str, clave: str, host: str = "", puerto: str = "", imap: str = "",
                          de: str = "", tls: str = "") -> Path:
    r = _config_env_ruta()
    r.parent.mkdir(parents=True, exist_ok=True)
    nuevas = {"CORREO_USUARIO": usuario, "CORREO_CLAVE": clave}
    if host:
        nuevas["CORREO_SMTP_HOST"] = host
    if puerto:
        nuevas["CORREO_SMTP_PORT"] = str(puerto)
    if imap:
        nuevas["CORREO_IMAP_HOST"] = imap
    if de:
        nuevas["CORREO_DE"] = de
    if tls:
        nuevas["CORREO_TLS"] = tls
    lineas = r.read_text(encoding="utf-8", errors="replace").splitlines() if r.exists() else []
    fuera = [l for l in lineas if not any(l.strip().startswith(k + "=") for k in nuevas)]
    fuera += ["%s=%s" % (k, v) for k, v in nuevas.items()]
    tmp = r.with_suffix(".env.tmp")
    tmp.write_text("\n".join(fuera) + "\n", encoding="utf-8")
    os.replace(tmp, r)
    return r


# ---------------------------------------------------------------------------
# Backends de correo
# ---------------------------------------------------------------------------

def outlook_tiene_cuenta() -> bool:
    """Mira el REGISTRO (no COM): un Dispatch sin cuenta abre el asistente y bloquea."""
    if os.name != "nt":
        return False
    try:
        import winreg  # type: ignore
    except Exception:
        return False
    for ver in ("16.0", "15.0"):
        try:
            base = winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Office\%s\Outlook\Profiles" % ver)
        except OSError:
            continue
        try:
            i = 0
            while True:
                try:
                    perfil = winreg.EnumKey(base, i)
                except OSError:
                    break
                i += 1
                # Cada subclave es un SERVICIO del perfil; la libreta de
                # direcciones cuenta como uno aunque no haya correo (medido
                # 2026-09-08: un perfil virgen tiene 'Libreta de direcciones
                # de Outlook' y nada mas). Cuenta de correo = tiene 'Email'.
                try:
                    k = winreg.OpenKey(base, perfil + r"\9375CFF0413111d3B88A00104B2A6676")
                except OSError:
                    continue
                j = 0
                while True:
                    try:
                        sub = winreg.EnumKey(k, j)
                    except OSError:
                        break
                    j += 1
                    try:
                        sk = winreg.OpenKey(k, sub)
                        try:
                            v, _t = winreg.QueryValueEx(sk, "Email")
                        except OSError:
                            v = None
                        if v and "@" in str(v):
                            return True
                    except OSError:
                        continue
        finally:
            winreg.CloseKey(base)
    return False


_OUTLOOK_ENVIAR = r'''
import sys, json
d = json.loads(sys.stdin.read())
import win32com.client as w
o = w.Dispatch("Outlook.Application")
m = o.CreateItem(0)
m.To = d["para"]
if d.get("cc"): m.CC = d["cc"]
m.Subject = d["asunto"]
if d.get("html"): m.HTMLBody = d["cuerpo"]
else: m.Body = d["cuerpo"]
for a in d.get("adjuntos") or []: m.Attachments.Add(a)
if d.get("borrador"):
    m.Save(); print("OK borrador guardado")
else:
    m.Send(); print("OK enviado")
'''

_OUTLOOK_LEER = r'''
import sys, json
d = json.loads(sys.stdin.read())
import win32com.client as w
o = w.Dispatch("Outlook.Application")
ns = o.GetNamespace("MAPI")
items = ns.GetDefaultFolder(6).Items
items.Sort("[ReceivedTime]", True)
out = []
for it in items:
    try:
        if d.get("no_leidos") and not it.UnRead: continue
        txt = (it.Subject or "") + " " + (it.SenderName or "") + " " + (it.Body or "")[:2000]
        if d.get("buscar") and d["buscar"].lower() not in txt.lower(): continue
        out.append({"de": "%s <%s>" % (it.SenderName, getattr(it, "SenderEmailAddress", "")),
                    "asunto": it.Subject, "fecha": str(it.ReceivedTime)[:16], "leido": not it.UnRead,
                    "cuerpo": (it.Body or "").strip()[:300]})
    except Exception:
        continue
    if len(out) >= d.get("n", 10): break
print("OK" + json.dumps(out, ensure_ascii=False))
'''

_OUTLOOK_CITA = r'''
import sys, json, datetime
d = json.loads(sys.stdin.read())
import win32com.client as w
o = w.Dispatch("Outlook.Application")
c = o.CreateItem(1)
c.Subject = d["asunto"]
c.Start = d["inicio"]
c.Duration = int(d.get("duracion", 60))
if d.get("lugar"): c.Location = d["lugar"]
if d.get("cuerpo"): c.Body = d["cuerpo"]
c.ReminderSet = True
c.Save()
print("OK cita guardada")
'''


def _outlook(script: str, datos: dict) -> tuple:
    """(ok, salida) corriendo el script COM en un subproceso con timeout."""
    try:
        r = subprocess.run([sys.executable, "-c", script], input=json.dumps(datos), capture_output=True,
                           text=True, encoding="utf-8", errors="replace", timeout=TIMEOUT_OUTLOOK_S)
    except subprocess.TimeoutExpired:
        return False, "Outlook no respondio en %ds (COM bloqueado: ¿asistente o dialogo abierto?)" % TIMEOUT_OUTLOOK_S
    out = (r.stdout or "").strip()
    if r.returncode == 0 and out.startswith("OK"):
        return True, out[2:].strip()
    return False, (r.stderr or out).strip()[-500:] or "exit %s" % r.returncode


def _smtp_enviar(c: dict, para: list, cc: list, asunto: str, cuerpo: str, adjuntos: list, html: bool) -> str:
    import smtplib
    from email.message import EmailMessage
    from email.utils import formatdate, make_msgid
    msg = EmailMessage()
    msg["From"] = c["de"] or c["usuario"]
    msg["To"] = ", ".join(para)
    if cc:
        msg["Cc"] = ", ".join(cc)
    msg["Subject"] = asunto
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid()
    if html:
        msg.set_content(re.sub(r"<[^>]+>", "", cuerpo))
        msg.add_alternative(cuerpo, subtype="html")
    else:
        msg.set_content(cuerpo)
    for a in adjuntos:
        p = Path(a)
        import mimetypes
        tipo, _ = mimetypes.guess_type(p.name)
        principal, sub = (tipo or "application/octet-stream").split("/", 1)
        msg.add_attachment(p.read_bytes(), maintype=principal, subtype=sub, filename=p.name)
    puerto = int(c.get("puerto") or 587)
    tls = str(c.get("tls", "1")).lower() in ("1", "on", "true", "yes", "si")
    if puerto == 465:
        s = smtplib.SMTP_SSL(c["host"], puerto, timeout=40)
    else:
        s = smtplib.SMTP(c["host"], puerto, timeout=40)
    try:
        s.ehlo()
        if tls and puerto != 465:
            s.starttls()
            s.ehlo()
        if c.get("usuario") and c.get("clave"):
            s.login(c["usuario"], c["clave"])
        s.send_message(msg, to_addrs=para + cc)
    finally:
        try:
            s.quit()
        except Exception:
            pass
    return "enviado por SMTP (%s:%d) a %s%s" % (c["host"], puerto, ", ".join(para), (" cc " + ", ".join(cc)) if cc else "")


def _partir_direcciones(s: str) -> list:
    return [x.strip() for x in re.split(r"[;,]\s*", (s or "").strip()) if x.strip()]


def enviar_correo(para: str, asunto: str, cuerpo: str, cc: str = "", adjunto: str = "", html: bool = False,
                  borrador: bool = False, ctx=None) -> str:
    destinos = _partir_direcciones(para)
    if not destinos or any("@" not in d for d in destinos):
        raise ValueError("destinatario invalido: %r (una o varias direcciones separadas por coma)" % para)
    if not asunto.strip():
        raise ValueError("falta asunto=")
    if not cuerpo.strip():
        raise ValueError("falta cuerpo=")
    adjuntos = []
    for a in _partir_direcciones(adjunto) if adjunto else []:
        adjuntos.append(str(PC.resolver_ruta(a, debe_existir=True, ctx=ctx)))
    c = config_correo()
    intentos = []
    if c["host"] and (c["usuario"] or c["host"] in ("127.0.0.1", "localhost")):
        if borrador:
            return "SMTP no guarda borradores: quita borrador=1 o usa Outlook con cuenta"
        try:
            return _smtp_enviar(c, destinos, _partir_direcciones(cc), asunto, cuerpo, adjuntos, html)
        except Exception as exc:
            intentos.append("SMTP %s: %s: %s" % (c["host"], type(exc).__name__, str(exc)[:200]))
    if outlook_tiene_cuenta():
        ok, salida = _outlook(_OUTLOOK_ENVIAR, {"para": "; ".join(destinos), "cc": cc, "asunto": asunto,
                                                "cuerpo": cuerpo, "adjuntos": adjuntos, "html": html,
                                                "borrador": borrador})
        if ok:
            return "%s por Outlook a %s" % (salida, ", ".join(destinos))
        intentos.append("Outlook: " + salida)
    else:
        intentos.append("Outlook: sin cuenta configurada (no se abre COM: bloquearia)")
    if not c["host"]:
        intentos.append("SMTP: sin configurar. Para Gmail: crea una clave de aplicacion en "
                        "https://myaccount.google.com/apppasswords y luego "
                        "`correo_configurar usuario=tu@gmail.com | clave=xxxx xxxx xxxx xxxx`")
    raise ValueError("no se pudo mandar: " + " · ".join(intentos))


def leer_correo(n: int = 10, buscar: str = "", no_leidos: bool = False) -> list:
    c = config_correo()
    if c.get("imap") and c.get("usuario") and c.get("clave"):
        import imaplib
        import email
        from email.header import decode_header, make_header
        m = imaplib.IMAP4_SSL(c["imap"], 993)
        try:
            m.login(c["usuario"], c["clave"])
            m.select("INBOX")
            crit = "UNSEEN" if no_leidos else "ALL"
            if buscar:
                crit = '(%s TEXT "%s")' % (crit, buscar.replace('"', ""))
            _st, datos = m.search(None, crit)
            ids = (datos[0].split() if datos and datos[0] else [])[-n:]
            out = []
            for i in reversed(ids):
                _st, partes = m.fetch(i, "(BODY.PEEK[])")
                crudo = partes[0][1] if partes and partes[0] else b""
                msg = email.message_from_bytes(crudo)

                def h(k):
                    try:
                        return str(make_header(decode_header(msg.get(k, "") or "")))
                    except Exception:
                        return msg.get(k, "") or ""
                cuerpo = ""
                for parte in msg.walk():
                    if parte.get_content_type() == "text/plain":
                        try:
                            cuerpo = parte.get_payload(decode=True).decode(parte.get_content_charset() or "utf-8", "replace")
                        except Exception:
                            cuerpo = ""
                        break
                out.append({"de": h("From"), "asunto": h("Subject"), "fecha": h("Date")[:25],
                            "cuerpo": " ".join(cuerpo.split())[:300]})
            return out
        finally:
            try:
                m.logout()
            except Exception:
                pass
    if outlook_tiene_cuenta():
        ok, salida = _outlook(_OUTLOOK_LEER, {"n": n, "buscar": buscar, "no_leidos": no_leidos})
        if ok:
            return json.loads(salida)
        raise ValueError("Outlook: " + salida)
    raise ValueError("no hay cuenta para leer correo: configura IMAP con `correo_configurar usuario=... | clave=...` "
                     "(Gmail: clave de aplicacion) o una cuenta en Outlook")


# ---------------------------------------------------------------------------
# Documentos
# ---------------------------------------------------------------------------

def escribir_documento(ruta: Path, texto: str, titulo: str = "") -> dict:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ext = ruta.suffix.lower()
    if ext == ".docx":
        try:
            import docx  # type: ignore
        except Exception as exc:
            raise ValueError("falta python-docx (%s). Instalalo con: pip install python-docx" % exc)
        d = docx.Document()
        if titulo:
            d.add_heading(titulo, level=1)
        parrafos = 0
        for bloque in re.split(r"\n\s*\n", texto.strip()):
            lineas = [l for l in bloque.splitlines() if l.strip()]
            if not lineas:
                continue
            if all(re.match(r"^\s*[-*•]\s+", l) for l in lineas):
                for l in lineas:
                    d.add_paragraph(re.sub(r"^\s*[-*•]\s+", "", l), style="List Bullet")
                    parrafos += 1
            elif len(lineas) == 1 and lineas[0].startswith("#"):
                d.add_heading(lineas[0].lstrip("#").strip(), level=min(3, max(1, len(lineas[0]) - len(lineas[0].lstrip("#")))))
            else:
                d.add_paragraph("\n".join(lineas))
                parrafos += 1
        d.save(str(ruta))
        return {"ruta": str(ruta), "formato": "docx", "parrafos": parrafos, "bytes": ruta.stat().st_size}
    contenido = texto if not titulo else ("%s\n%s\n\n%s" % (titulo, "=" * len(titulo), texto) if ext != ".md"
                                          else "# %s\n\n%s" % (titulo, texto))
    ruta.write_text(contenido.rstrip("\n") + "\n", encoding="utf-8")
    return {"ruta": str(ruta), "formato": ext.lstrip(".") or "txt", "parrafos": len(re.split(r"\n\s*\n", contenido.strip())),
            "bytes": ruta.stat().st_size}


# ---------------------------------------------------------------------------
# Abrir cosas en el escritorio de Cognia
# ---------------------------------------------------------------------------

def _navegador() -> list:
    for exe in (r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
                r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if Path(exe).exists():
            perfil = Path.home() / ".cognia" / "cotidiano" / ("perfil_" + Path(exe).stem)
            perfil.mkdir(parents=True, exist_ok=True)
            return [exe, "--new-window", "--no-first-run", "--no-default-browser-check",
                    "--user-data-dir=%s" % perfil]
    return []


def abrir_en_escritorio(objetivo: str, titulo: str = "", espera_ms: int = ESPERA_VENTANA_MS, ctx=None) -> dict:
    """Abre un fichero, una URL o una app y muda la ventana al escritorio de Cognia.
    Devuelve {hwnd, titulo, mudada, app_id, pid, via}."""
    from cognia.agent import escritorio_propio as EP
    from cognia.agent import app_tools as AT
    o = (objetivo or "").strip().strip("\"'")
    if not o:
        raise ValueError("falta que abrir (ruta, URL o app)")
    es_url = bool(re.match(r"^(https?://|www\.)", o, re.I))
    if es_url:
        if o.lower().startswith("www."):
            o = "https://" + o
        nav = _navegador()
        if not nav:
            raise ValueError("no encontre Chrome ni Edge para abrir la URL")
        cmd = " ".join('"%s"' % p if " " in p else p for p in nav + [o])
        app_id, a = AT.lanzar(cmd, espera_ms=max(espera_ms, 6000), titulo=titulo)
        return {"hwnd": a["hwnd"], "titulo": a["titulo"], "mudada": a["mudada"], "app_id": app_id, "pid": a["pid"], "via": "navegador"}
    ruta = None
    try:
        ruta = PC.resolver_ruta(o, debe_existir=True, ctx=ctx)
    except Exception:
        ruta = None
    if ruta is not None and ruta.is_file():
        if ruta.suffix.lower() in (".txt", ".md", ".log", ".csv", ".json", ".py", ".ini", ".yaml", ".yml"):
            app_id, a = AT.lanzar('notepad.exe "%s"' % ruta, espera_ms=espera_ms, titulo=titulo or ruta.name)
            return {"hwnd": a["hwnd"], "titulo": a["titulo"], "mudada": a["mudada"], "app_id": app_id, "pid": a["pid"], "via": "notepad"}
        # asociacion de Windows (Word, Excel, visor de fotos...): no es hijo nuestro,
        # asi que se detecta por DIFERENCIA de ventanas visibles antes/despues.
        antes = {h for h, _p, _t in EP.ventanas_visibles()}
        os.startfile(str(ruta))  # noqa: S606
        hwnd = _esperar_ventana_nueva(antes, titulo or ruta.stem, espera_ms)
        if hwnd is None:
            raise ValueError("abri %s pero no aparecio ninguna ventana nueva en %.1fs" % (ruta.name, espera_ms / 1000))
        mudada = EP.mover_ventana(hwnd) if EP.activo() else False
        app_id = _registrar_ventana(hwnd, str(ruta), mudada)
        return {"hwnd": hwnd, "titulo": EP.titulo_ventana(hwnd), "mudada": mudada, "app_id": app_id,
                "pid": EP.pid_de(hwnd), "via": "asociacion"}
    # una app o un comando
    app_id, a = AT.lanzar(o, espera_ms=espera_ms, titulo=titulo)
    return {"hwnd": a["hwnd"], "titulo": a["titulo"], "mudada": a["mudada"], "app_id": app_id, "pid": a["pid"], "via": "app"}


def _esperar_ventana_nueva(antes: set, pista: str, espera_ms: int):
    """La ventana nueva cuyo titulo lleva la pista (el nombre del fichero). Se
    ESPERA a la pista hasta agotar el tiempo: Word abre primero un splash
    'Abriendo - Word' que muere enseguida, y quedarse con la primera ventana
    nueva era mudar un fantasma (medido 2026-09-08). Sin pista al final, la
    primera nueva que siga viva."""
    from cognia.agent import escritorio_propio as EP
    t0 = time.time()
    pista = (pista or "").lower()
    primera = None
    while time.time() - t0 < max(1.0, espera_ms / 1000.0):
        time.sleep(0.3)
        nuevas = [(h, p, t) for h, p, t in EP.ventanas_visibles()
                  if h not in antes and (t or "").strip() and EP.rect_ventana(h)[2] > 100]
        if not nuevas:
            continue
        con_pista = [x for x in nuevas if pista and pista in (x[2] or "").lower()]
        if con_pista:
            return con_pista[0][0]
        if primera is None or not EP.ventana_viva(primera):
            primera = nuevas[0][0]
        if not pista:
            return primera
    return primera if (primera is not None and EP.ventana_viva(primera)) else None


def _registrar_ventana(hwnd: int, origen: str, mudada: bool) -> str:
    """La deja en el registro de app_tools para app_ver / app_teclas / app_cerrar."""
    from cognia.agent import app_tools as AT
    from cognia.agent import escritorio_propio as EP
    try:
        app_id = AT._nuevo_id()
        AT._APPS[app_id] = {"pid": EP.pid_de(hwnd), "hwnd": hwnd, "cmd": origen, "titulo": EP.titulo_ventana(hwnd),
                            "ts": time.time(), "log": "", "proc": None, "capturas": [], "mudada": mudada,
                            "escritorio": EP.config()["nombre"] if mudada else "actual"}
        AT._persistir()
        return app_id
    except Exception as exc:
        _avisar("no se pudo registrar la ventana en app_tools: %s" % exc)
        return ""


def _captura(hwnd: int, ctx, nombre: str) -> str:
    from cognia.agent import escritorio_propio as EP
    try:
        png = PC.ruta_salida(ctx, nombre, ".png")
        r = EP.capturar_ventana(hwnd, png)
        res = PC.resumen_imagen(png)
        return "%s · %s" % (r["png"], PC.texto_resumen_imagen(res))
    except Exception as exc:
        return "sin captura: %s" % exc


# ---------------------------------------------------------------------------
# Calendario y recordatorios
# ---------------------------------------------------------------------------

def _parsear_fecha(s: str) -> _dt.datetime:
    s = (s or "").strip()
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M", "%d/%m/%Y %H:%M", "%Y-%m-%d"):
        try:
            return _dt.datetime.strptime(s, fmt)
        except ValueError:
            continue
    m = re.match(r"^(hoy|manana|mañana)\s+(\d{1,2}):(\d{2})$", s.lower())
    if m:
        base = _dt.datetime.now().replace(hour=int(m.group(2)), minute=int(m.group(3)), second=0, microsecond=0)
        return base + _dt.timedelta(days=0 if m.group(1) == "hoy" else 1)
    raise ValueError("fecha no entendida: %r (usa YYYY-MM-DD HH:MM, hoy HH:MM o manana HH:MM)" % s)


def agregar_cita(asunto: str, inicio: str, duracion: int = 60, lugar: str = "", cuerpo: str = "") -> dict:
    t = _parsear_fecha(inicio)
    if outlook_tiene_cuenta():
        ok, salida = _outlook(_OUTLOOK_CITA, {"asunto": asunto, "inicio": t.strftime("%Y-%m-%d %H:%M"),
                                              "duracion": duracion, "lugar": lugar, "cuerpo": cuerpo})
        if ok:
            return {"via": "outlook", "detalle": salida, "inicio": t.isoformat(timespec="minutes")}
        _avisar("Outlook cita: %s" % salida)
    carpeta = Path.home() / ".cognia" / "cotidiano" / "citas"
    carpeta.mkdir(parents=True, exist_ok=True)
    fin = t + _dt.timedelta(minutes=int(duracion))
    nombre = re.sub(r"[^A-Za-z0-9_-]+", "_", asunto)[:40] or "cita"
    ics = carpeta / ("%s_%s.ics" % (t.strftime("%Y%m%d_%H%M"), nombre))
    ics.write_text("BEGIN:VCALENDAR\r\nVERSION:2.0\r\nPRODID:-//Cognia//cotidiano//ES\r\nBEGIN:VEVENT\r\n"
                   "UID:%s@cognia\r\nDTSTAMP:%s\r\nDTSTART:%s\r\nDTEND:%s\r\nSUMMARY:%s\r\n%s%s"
                   "END:VEVENT\r\nEND:VCALENDAR\r\n"
                   % (ics.stem, _dt.datetime.utcnow().strftime("%Y%m%dT%H%M%SZ"), t.strftime("%Y%m%dT%H%M%S"),
                      fin.strftime("%Y%m%dT%H%M%S"), asunto, ("LOCATION:%s\r\n" % lugar) if lugar else "",
                      ("DESCRIPTION:%s\r\n" % cuerpo.replace("\n", "\\n")) if cuerpo else ""), encoding="utf-8")
    return {"via": "ics", "detalle": str(ics), "inicio": t.isoformat(timespec="minutes")}


def _cuando(en: str = "", a: str = "") -> _dt.datetime:
    ahora = _dt.datetime.now()
    if en:
        m = re.match(r"^(\d+)\s*(m|min|minutos?|h|horas?|s|seg)?$", en.strip().lower())
        if not m:
            raise ValueError("en= admite 20m, 2h o 90s")
        n, u = int(m.group(1)), (m.group(2) or "m")[0]
        return ahora + _dt.timedelta(seconds=n * {"m": 60, "h": 3600, "s": 1}[u])
    if a:
        m = re.match(r"^(\d{1,2}):(\d{2})$", a.strip())
        if not m:
            return _parsear_fecha(a)
        t = ahora.replace(hour=int(m.group(1)), minute=int(m.group(2)), second=0, microsecond=0)
        return t if t > ahora else t + _dt.timedelta(days=1)
    raise ValueError("di cuando: en=20m o a=HH:MM")


def programar_recordatorio(texto: str, cuando: _dt.datetime) -> dict:
    """Tarea programada de Windows que muestra el aviso con msg.exe (sesion del dueno)."""
    if os.name != "nt":
        raise ValueError("los recordatorios solo existen en Windows (schtasks)")
    texto = " ".join(texto.split())
    if not texto:
        raise ValueError("falta el texto del recordatorio")
    if cuando <= _dt.datetime.now():
        raise ValueError("la hora ya paso")
    nombre = "Cognia_recordatorio_%s" % cuando.strftime("%Y%m%d_%H%M%S")
    seguro = texto.replace('"', "'")[:200]
    tr = 'msg * /TIME:600 "Cognia: %s"' % seguro
    r = subprocess.run(["schtasks", "/Create", "/F", "/SC", "ONCE", "/TN", nombre, "/TR", tr,
                        "/ST", cuando.strftime("%H:%M"), "/SD", cuando.strftime("%d/%m/%Y")],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30)
    if r.returncode != 0:
        # el formato de fecha depende del locale: segundo intento con MM/DD/YYYY
        r = subprocess.run(["schtasks", "/Create", "/F", "/SC", "ONCE", "/TN", nombre, "/TR", tr,
                            "/ST", cuando.strftime("%H:%M"), "/SD", cuando.strftime("%m/%d/%Y")],
                           capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30)
    if r.returncode != 0:
        raise ValueError("schtasks fallo: %s" % (r.stderr or r.stdout).strip()[-300:])
    return {"tarea": nombre, "cuando": cuando.strftime("%Y-%m-%d %H:%M"), "texto": texto}


def listar_recordatorios() -> list:
    if os.name != "nt":
        return []
    r = subprocess.run(["schtasks", "/Query", "/FO", "CSV", "/NH"], capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=30)
    out = []
    for linea in (r.stdout or "").splitlines():
        if "Cognia_recordatorio_" in linea:
            partes = [p.strip('"') for p in linea.split('","')]
            out.append({"tarea": partes[0].strip('"').lstrip("\\"), "proxima": partes[1] if len(partes) > 1 else ""})
    return out


def borrar_recordatorio(nombre: str) -> bool:
    r = subprocess.run(["schtasks", "/Delete", "/F", "/TN", nombre], capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=30)
    return r.returncode == 0


# ---------------------------------------------------------------------------
# Estado
# ---------------------------------------------------------------------------

def estado() -> dict:
    c = config_correo()
    try:
        import docx  # type: ignore # noqa: F401
        word = "python-docx"
    except Exception:
        word = "sin python-docx (pip install python-docx)"
    try:
        from cognia.agent import escritorio_propio as EP
        e = EP.estado()
        esc = "activo (n%s)" % e.get("numero") if e.get("activo") else ("apagado" if not e.get("configurado") else "no disponible: " + str(e.get("motivo", "")))
    except Exception as exc:
        esc = "error: %s" % exc
    return {
        "correo_smtp": ("%s:%s como %s" % (c["host"], c["puerto"], c["usuario"] or "(sin usuario)")) if c["host"] else "sin configurar",
        "correo_imap": c["imap"] if (c["imap"] and c["usuario"]) else "sin configurar",
        "outlook": "con cuenta" if outlook_tiene_cuenta() else "sin cuenta (no se usa COM)",
        "documentos": word, "navegador": (_navegador() or ["ninguno"])[0], "escritorio": esc,
        "recordatorios": len(listar_recordatorios()) if os.name == "nt" else 0, "ultimo": dict(_ULTIMO),
    }


def texto_estado() -> str:
    e = estado()
    lineas = ["correo (enviar): SMTP %s · Outlook %s" % (e["correo_smtp"], e["outlook"]),
              "correo (leer): IMAP %s" % e["correo_imap"],
              "documentos: %s · navegador: %s" % (e["documentos"], e["navegador"]),
              "escritorio propio: %s · recordatorios programados: %d" % (e["escritorio"], e["recordatorios"])]
    u = e["ultimo"]
    if u.get("tool"):
        lineas.append("ultimo: %s %s%s" % (u["tool"], u.get("detalle", ""), (" · error: " + u["error"]) if u.get("error") else ""))
    return "\n".join(lineas)


# ---------------------------------------------------------------------------
# Registro de tools
# ---------------------------------------------------------------------------

def register(tool) -> None:
    @tool("documento_escribir",
          "documento_escribir <ruta.docx|.txt|.md> | texto=... [| titulo=...] [| abrir=0]"
          "  -- escribe un documento (Word real si es .docx) y lo abre en el escritorio propio de Cognia",
          desc="Escribe un documento con el texto dado: .docx de Word (titulo como encabezado, parrafos, "
               "lineas con '- ' como vinetas, '# ' como subtitulos), o .txt/.md tal cual. Por defecto lo ABRE "
               "en el escritorio propio de Cognia (Word o Bloc de notas) y devuelve la ruta, la ventana y una "
               "captura; abrir=0 solo escribe. Usala cuando el usuario pida 'escribeme una carta / un informe / "
               "una nota' o quiera ver el texto en un documento.",
          params=[{"nombre": "ruta", "tipo": "string", "requerido": True, "descripcion": "ruta del documento (.docx, .txt, .md)"},
                  {"nombre": "texto", "tipo": "string", "requerido": True, "clave": True, "descripcion": "el contenido completo"},
                  {"nombre": "titulo", "tipo": "string", "requerido": False, "clave": True, "descripcion": "titulo/encabezado"},
                  {"nombre": "abrir", "tipo": "integer", "requerido": False, "clave": True, "descripcion": "1 (default) abre la ventana en el escritorio de Cognia; 0 no"}],
          danger=False, timeout_s=90)
    def _documento_escribir(args, ctx):
        ruta, o = PC.partir_args(args, _CLAVES)
        try:
            if not ruta:
                raise ValueError("falta la ruta del documento")
            texto = o.get("texto", "")
            if not texto.strip():
                raise ValueError("falta texto= (el contenido)")
            # Un documento NUEVO va al workspace del usuario (no al scratchpad
            # de la tarea, que es donde resolver_ruta deja lo que no existe).
            p = Path(ruta.strip("\"'"))
            if not p.is_absolute():
                base = (ctx or {}).get("workspace") if isinstance(ctx, dict) else None
                p = Path(str(base)) / p if base else PC.bases_relativas()[0] / p
            if p.suffix.lower() not in (".docx", ".txt", ".md"):
                p = p.with_suffix(".docx")
            r = escribir_documento(p, texto.replace("\\n", "\n"), o.get("titulo", ""))
            msg = "RESULTADO documento_escribir: %s (%s, %d parrafos, %d bytes)" % (r["ruta"], r["formato"], r["parrafos"], r["bytes"])
            if PC.entero(o.get("abrir"), 1, 0, 1):
                try:
                    v = abrir_en_escritorio(r["ruta"], espera_ms=ESPERA_VENTANA_MS, ctx=ctx)
                    msg += " · abierto en %s (ventana %r%s%s) · captura %s" % (
                        "el escritorio de Cognia" if v["mudada"] else "el escritorio actual", v["titulo"],
                        (", app " + v["app_id"]) if v.get("app_id") else "", "", _captura(v["hwnd"], ctx, "documento_" + p.stem))
                except Exception as exc:
                    msg += " · escrito, pero no se pudo abrir: %s" % exc
            _anotar("documento_escribir", r["ruta"])
            return msg
        except Exception as exc:
            return _err("documento_escribir", exc)

    @tool("correo_enviar",
          "correo_enviar <para> | asunto=... | cuerpo=... [| cc=...] [| adjunto=ruta] [| html=1]"
          "  -- manda un correo (SMTP configurado u Outlook con cuenta)",
          desc="Manda un correo electronico a una o varias direcciones (separadas por coma) con asunto, cuerpo "
               "(texto; html=1 si el cuerpo es HTML), copia y adjuntos. Usa el SMTP configurado (Gmail con clave "
               "de aplicacion: `correo_configurar`) o Outlook si tiene cuenta. Si no hay ninguno, te dice "
               "exactamente que falta: NO inventes que lo mandaste. Usala cuando el usuario pida 'mandale un "
               "correo a X' o 'envia este informe por mail'.",
          params=[{"nombre": "para", "tipo": "string", "requerido": True, "descripcion": "destinatario(s)"},
                  {"nombre": "asunto", "tipo": "string", "requerido": True, "clave": True, "descripcion": "asunto"},
                  {"nombre": "cuerpo", "tipo": "string", "requerido": True, "clave": True, "descripcion": "texto del correo"},
                  {"nombre": "cc", "tipo": "string", "requerido": False, "clave": True, "descripcion": "copia"},
                  {"nombre": "adjunto", "tipo": "string", "requerido": False, "clave": True, "descripcion": "ruta(s) a adjuntar, separadas por coma"},
                  {"nombre": "html", "tipo": "integer", "requerido": False, "clave": True, "descripcion": "1 si el cuerpo es HTML"}],
          danger=True, timeout_s=120)
    def _correo_enviar(args, ctx):
        para, o = PC.partir_args(args, _CLAVES)
        try:
            r = enviar_correo(para, o.get("asunto", ""), o.get("cuerpo", "").replace("\\n", "\n"), cc=o.get("cc", ""),
                              adjunto=o.get("adjunto", ""), html=bool(PC.entero(o.get("html"), 0, 0, 1)), ctx=ctx)
            _anotar("correo_enviar", "%s: %s" % (para, o.get("asunto", "")))
            return "RESULTADO correo_enviar: %s · asunto %r" % (r, o.get("asunto", ""))
        except Exception as exc:
            return _err("correo_enviar", exc)

    @tool("correo_leer",
          "correo_leer [| n=10] [| buscar=texto] [| no_leidos=1]  -- lee los ultimos correos de la bandeja (IMAP u Outlook)",
          desc="Lista los ultimos correos recibidos (remitente, asunto, fecha y el principio del cuerpo), "
               "filtrando por texto o solo no leidos. Usa IMAP con la cuenta configurada o Outlook con cuenta.",
          params=[{"nombre": "n", "tipo": "integer", "requerido": False, "clave": True, "descripcion": "cuantos (default 10)"},
                  {"nombre": "buscar", "tipo": "string", "requerido": False, "clave": True, "descripcion": "texto a buscar"},
                  {"nombre": "no_leidos", "tipo": "integer", "requerido": False, "clave": True, "descripcion": "1 = solo no leidos"}],
          danger=False, timeout_s=120)
    def _correo_leer(args, ctx):
        _s, o = PC.partir_args(args, _CLAVES)
        try:
            items = leer_correo(PC.entero(o.get("n"), 10, 1, 50), o.get("buscar", ""), bool(PC.entero(o.get("no_leidos"), 0, 0, 1)))
            if not items:
                return "RESULTADO correo_leer: ningun correo que case"
            lineas = ["%d. %s · %s · %s\n   %s" % (i, it.get("fecha", ""), it.get("de", ""), it.get("asunto", ""), it.get("cuerpo", ""))
                      for i, it in enumerate(items, 1)]
            _anotar("correo_leer", "%d correos" % len(items))
            return "RESULTADO correo_leer (%d):\n%s" % (len(items), "\n".join(lineas))
        except Exception as exc:
            return _err("correo_leer", exc)

    @tool("correo_configurar",
          "correo_configurar usuario=tu@gmail.com | clave=xxxx [| host=smtp.gmail.com] [| puerto=587] [| imap=imap.gmail.com]"
          "  -- guarda la cuenta de correo (clave de aplicacion) en ~/.cognia/config.env",
          desc="Configura la cuenta con la que Cognia manda y lee correo. Para Gmail hay que crear una clave de "
               "aplicacion (myaccount.google.com/apppasswords). La clave se guarda en ~/.cognia/config.env, "
               "nunca en el historial. Pidele la clave al usuario; no la inventes.",
          params=[{"nombre": "usuario", "tipo": "string", "requerido": True, "clave": True, "descripcion": "direccion de correo"},
                  {"nombre": "clave", "tipo": "string", "requerido": True, "clave": True, "descripcion": "clave de aplicacion"},
                  {"nombre": "host", "tipo": "string", "requerido": False, "clave": True, "descripcion": "servidor SMTP"},
                  {"nombre": "puerto", "tipo": "integer", "requerido": False, "clave": True, "descripcion": "puerto SMTP (587 o 465)"},
                  {"nombre": "imap", "tipo": "string", "requerido": False, "clave": True, "descripcion": "servidor IMAP"}],
          danger=True, timeout_s=30)
    def _correo_configurar(args, ctx):
        _s, o = PC.partir_args(args, _CLAVES)
        try:
            usuario, clave = o.get("usuario", "").strip(), o.get("clave", "").strip()
            if "@" not in usuario or not clave:
                raise ValueError("faltan usuario= y clave=")
            r = guardar_config_correo(usuario, clave, o.get("host", ""), o.get("puerto", ""), o.get("imap", ""),
                                      o.get("de", ""), o.get("tls", ""))
            c = config_correo()
            _anotar("correo_configurar", usuario)
            return ("RESULTADO correo_configurar: guardado en %s · SMTP %s:%s · IMAP %s · usuario %s (la clave no se muestra)"
                    % (r, c["host"] or "?", c["puerto"], c["imap"] or "?", usuario))
        except Exception as exc:
            return _err("correo_configurar", exc)

    @tool("calendario_agregar",
          "calendario_agregar <asunto> | inicio=YYYY-MM-DD HH:MM [| duracion=60] [| lugar=...] [| cuerpo=...]"
          "  -- crea una cita (Outlook con cuenta, o un .ics importable)",
          desc="Agenda una cita: en el calendario de Outlook si hay cuenta; si no, deja un fichero .ics en "
               "~/.cognia/cotidiano/citas listo para importar. inicio admite 'hoy 16:00' y 'manana 09:30'.",
          params=[{"nombre": "asunto", "tipo": "string", "requerido": True, "descripcion": "titulo de la cita"},
                  {"nombre": "inicio", "tipo": "string", "requerido": True, "clave": True, "descripcion": "YYYY-MM-DD HH:MM | hoy HH:MM | manana HH:MM"},
                  {"nombre": "duracion", "tipo": "integer", "requerido": False, "clave": True, "descripcion": "minutos (default 60)"},
                  {"nombre": "lugar", "tipo": "string", "requerido": False, "clave": True, "descripcion": "lugar"},
                  {"nombre": "cuerpo", "tipo": "string", "requerido": False, "clave": True, "descripcion": "notas"}],
          danger=False, timeout_s=90)
    def _calendario_agregar(args, ctx):
        asunto, o = PC.partir_args(args, _CLAVES)
        try:
            if not asunto.strip():
                raise ValueError("falta el asunto")
            r = agregar_cita(asunto.strip(), o.get("inicio", ""), PC.entero(o.get("duracion"), 60, 5, 24 * 60),
                             o.get("lugar", ""), o.get("cuerpo", ""))
            _anotar("calendario_agregar", "%s %s" % (asunto, r["inicio"]))
            return "RESULTADO calendario_agregar: %r el %s · %s" % (asunto.strip(), r["inicio"],
                                                                    ("Outlook: " + r["detalle"]) if r["via"] == "outlook" else "guardada como " + r["detalle"] + " (Outlook sin cuenta)")
        except Exception as exc:
            return _err("calendario_agregar", exc)

    @tool("recordatorio",
          "recordatorio <texto> | en=20m | a=HH:MM  -- programa un aviso en pantalla a esa hora (o 'lista' / 'borrar <tarea>')",
          desc="Programa un recordatorio que aparece en la pantalla del usuario a la hora dicha (tarea programada de "
               "Windows con msg.exe). en=20m / en=2h para relativo, a=HH:MM para una hora. `recordatorio lista` "
               "muestra los pendientes y `recordatorio borrar <tarea>` quita uno.",
          params=[{"nombre": "texto", "tipo": "string", "requerido": True, "descripcion": "que recordar, o 'lista' / 'borrar <tarea>'"},
                  {"nombre": "en", "tipo": "string", "requerido": False, "clave": True, "descripcion": "20m, 2h, 90s"},
                  {"nombre": "a", "tipo": "string", "requerido": False, "clave": True, "descripcion": "HH:MM (hoy, o manana si ya paso)"}],
          danger=True, timeout_s=60)
    def _recordatorio(args, ctx):
        texto, o = PC.partir_args(args, _CLAVES)
        try:
            bajo = texto.strip().lower()
            if bajo == "lista":
                rs = listar_recordatorios()
                if not rs:
                    return "RESULTADO recordatorio lista: ninguno programado"
                return "RESULTADO recordatorio lista:\n" + "\n".join("  %s · proxima %s" % (r["tarea"], r["proxima"]) for r in rs)
            if bajo.startswith("borrar "):
                nombre = texto.strip().split(None, 1)[1].strip()
                return "RESULTADO recordatorio borrar: %s" % ("borrado " + nombre if borrar_recordatorio(nombre) else "ERROR no se pudo borrar " + nombre)
            r = programar_recordatorio(texto, _cuando(o.get("en", ""), o.get("a", "")))
            _anotar("recordatorio", "%s @ %s" % (r["texto"], r["cuando"]))
            return "RESULTADO recordatorio: %r programado para el %s (tarea %s)" % (r["texto"], r["cuando"], r["tarea"])
        except Exception as exc:
            return _err("recordatorio", exc)

    @tool("abrir_en_escritorio",
          "abrir_en_escritorio <ruta|URL|app> [| titulo=...] [| espera=MS]"
          "  -- abre un fichero, una pagina o una app y muda su ventana al escritorio propio de Cognia",
          desc="Abre lo que le des (un documento, una URL en una ventana nueva del navegador con perfil propio, "
               "un .exe o un comando con ventana) y lo lleva al escritorio virtual de Cognia sin molestar al "
               "usuario. Devuelve el id de app (para app_ver / app_teclas / app_clic / app_cerrar) y una captura.",
          params=[{"nombre": "objetivo", "tipo": "string", "requerido": True, "descripcion": "ruta, URL, .exe o comando"},
                  {"nombre": "titulo", "tipo": "string", "requerido": False, "clave": True, "descripcion": "parte del titulo de la ventana esperada"},
                  {"nombre": "espera", "tipo": "integer", "requerido": False, "clave": True, "descripcion": "ms maximos hasta que aparezca la ventana"}],
          danger=True, timeout_s=90)
    def _abrir(args, ctx):
        objetivo, o = PC.partir_args(args, _CLAVES)
        try:
            v = abrir_en_escritorio(objetivo, o.get("titulo", ""), PC.entero(o.get("espera"), ESPERA_VENTANA_MS, 500, 60000), ctx=ctx)
            _anotar("abrir_en_escritorio", "%s -> %r" % (objetivo, v["titulo"]))
            return ("RESULTADO abrir_en_escritorio: ventana %r (hwnd %d, via %s) en %s%s · captura %s\n"
                    "Siguientes: app_ver %s · app_teclas %s | escribir \"hola\" · app_cerrar %s"
                    % (v["titulo"], v["hwnd"], v["via"], "el escritorio de Cognia" if v["mudada"] else "el escritorio ACTUAL (escritorio propio apagado o no disponible)",
                       (" · app " + v["app_id"]) if v.get("app_id") else "", _captura(v["hwnd"], ctx, "abrir_" + re.sub(r"[^A-Za-z0-9]+", "_", objetivo)[:30]),
                       v.get("app_id", "?"), v.get("app_id", "?"), v.get("app_id", "?")))
        except Exception as exc:
            return _err("abrir_en_escritorio", exc)

    @tool("cotidiano_estado",
          "cotidiano_estado  -- que backend hay para correo, documentos, calendario y escritorio, y el ultimo error",
          desc="Diagnostico de las tareas cotidianas: SMTP/IMAP/Outlook configurados o no, python-docx, navegador, "
               "escritorio propio y recordatorios programados.",
          params=[], danger=False, timeout_s=60)
    def _estado(args, ctx):
        try:
            return "RESULTADO cotidiano_estado:\n" + texto_estado()
        except Exception as exc:
            return _err("cotidiano_estado", exc)
